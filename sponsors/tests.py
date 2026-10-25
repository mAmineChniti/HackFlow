"""
Tests unitaires et d'intégration pour le module Gestion des Sponsors.
"""

import io
import json
import os
from decimal import Decimal
from unittest.mock import MagicMock, patch
import openpyxl
from django.contrib import admin
from django.contrib.auth.models import User
from django.core import mail
from django.core.cache import cache
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from hackathons.models import Hackathon
from sponsors.ai import (
    generate_prospecting_email,
    get_hackathon_text,
    recommend_sponsors_for_hackathon,
)
from sponsors.ai.email_generator import build_sanitized_ai_payload
from sponsors.models import (
    Sponsor,
    SponsorEmailLog,
    Sponsorship,
    StatusChoices,
    StatusHistory,
    TierChoices,
)
from sponsors.permissions import (
    ROLE_PERMISSIONS,
    SPONSOR_PERMISSIONS,
    can_manage_sponsorship,
    can_view_confidential,
    get_user_roles,
    sponsorships_visible_to,
    user_has_sponsor_permission,
)
from sponsors.templatetags.sponsor_tags import get_confirmed_sponsors


class SponsorAdminInterfaceTests(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_superuser(
            username='sponsor_admin',
            email='admin@example.com',
            password='StrongPassword123!',
        )
        self.client.force_login(self.admin_user)
        self.sponsor = Sponsor.objects.create(name='Admin Template Sponsor')
        self.hackathon = Hackathon.objects.create(
            name='Admin Template Hackathon',
            theme='Technology',
            start_date=timezone.now(),
            end_date=timezone.now() + timezone.timedelta(days=1),
        )
        self.sponsorship = Sponsorship.objects.create(
            sponsor=self.sponsor,
            hackathon=self.hackathon,
        )

    def test_admin_root_responds_200(self):
        response = self.client.get(reverse('admin:index'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Django administration')
        self.assertContains(response, reverse('admin:sponsors_sponsor_changelist'))

    def test_sponsor_lists_do_not_expose_contact_emails_or_internal_notes(self):
        self.sponsor.contact_email = 'private-contact@example.com'
        self.sponsor.save(update_fields=('contact_email',))
        self.sponsorship.notes = 'Confidential negotiation note'
        self.sponsorship.save(update_fields=('notes',))

        sponsor_response = self.client.get(reverse('admin:sponsors_sponsor_changelist'))
        sponsorship_response = self.client.get(
            reverse('admin:sponsors_sponsorship_changelist')
        )

        self.assertNotContains(sponsor_response, 'private-contact@example.com')
        self.assertNotContains(sponsorship_response, 'private-contact@example.com')
        self.assertNotContains(sponsorship_response, 'Confidential negotiation note')

    def test_sponsorship_admin_actions_use_status_service_and_record_history(self):
        actions = (
            ('mark_contacted', StatusChoices.CONTACTE),
            ('mark_negotiation', StatusChoices.NEGOCIATION),
            ('mark_confirmed', StatusChoices.CONFIRME),
            ('mark_refused', StatusChoices.REFUSE),
        )
        url = reverse('admin:sponsors_sponsorship_changelist')

        for action, expected_status in actions:
            with self.subTest(action=action):
                response = self.client.post(
                    url,
                    {
                        'action': action,
                        '_selected_action': [str(self.sponsorship.pk)],
                        'index': 0,
                        'post': 'yes',
                    },
                )
                self.assertEqual(response.status_code, 302)
                self.sponsorship.refresh_from_db()
                self.assertEqual(self.sponsorship.status, expected_status)

        self.assertEqual(
            StatusHistory.objects.filter(sponsorship=self.sponsorship).count(),
            len(actions),
        )
        self.assertEqual(
            StatusHistory.objects.filter(sponsorship=self.sponsorship)
            .order_by('changed_at')
            .last()
            .new_status,
            StatusChoices.REFUSE,
        )

    def test_status_and_email_logs_are_read_only_in_admin(self):
        response = self.client.get(reverse('admin:sponsors_statushistory_changelist'))
        request = response.wsgi_request
        for model in (StatusHistory, SponsorEmailLog):
            model_admin = admin.site._registry[model]
            with self.subTest(model=model.__name__):
                self.assertFalse(model_admin.has_add_permission(request))
                self.assertFalse(model_admin.has_change_permission(request))
                self.assertFalse(model_admin.has_delete_permission(request))

    def test_non_confidential_roles_cannot_open_sponsor_admin(self):
        staff_user = User.objects.create_superuser(
            username='sponsor_admin_student_role',
            email='student-role@example.com',
            password='StrongPassword123!',
        )
        self.client.force_login(staff_user)

        with patch('sponsors.permissions.get_user_roles', return_value={'etudiant'}):
            response = self.client.get(reverse('admin:sponsors_sponsor_changelist'))

        self.assertEqual(response.status_code, 403)

class SponsorModelAndViewsTest(TestCase):
    def setUp(self):
        cache.clear()
        # Utilisateur normal (non-staff)
        self.regular_user = User.objects.create_user(
            username="student_user",
            password="password123",
            email="student@esprit.tn",
        )
        # Utilisateur staff / organisateur
        self.staff_user = User.objects.create_user(
            username="staff_user",
            password="password123",
            email="staff@esprit.tn",
            is_staff=True,
        )

        # Hackathon de test
        self.hackathon = Hackathon.objects.create(
            name="ESPRIT Hack 2026",
            theme="AI & Cloud Innovation",
            start_date=timezone.now(),
            end_date=timezone.now() + timezone.timedelta(days=2),
        )

        # Sponsor de test
        self.sponsor = Sponsor.objects.create(
            name="Alpha Corp",
            website="https://alpha.example.com",
            sector="tech",
            description="Leader des solutions cloud pour l'éducation.",
            contact_name="Karim Tounsi",
            contact_email="karim@alpha.example.com",
        )

        # Sponsoring de test
        self.sponsorship = Sponsorship.objects.create(
            sponsor=self.sponsor,
            hackathon=self.hackathon,
            tier=TierChoices.PLATINUM,
            amount=Decimal("5000.00"),
            status=StatusChoices.CONFIRME,
        )

    def test_public_view_accessible_by_anonymous(self):
        """La page publique doit être accessible sans authentification (200 OK)."""
        response = self.client.get(reverse("sponsors:public_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Alpha Corp")
        self.assertContains(response, "Platinum partners")

    def test_dashboard_requires_staff(self):
        """Le dashboard doit refuser les utilisateurs anonymes et non-staff."""
        # Anonyme -> Redirige vers login
        response = self.client.get(reverse("sponsors:dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith("/accounts/login/"))

        # Utilisateur non-staff -> Redirige vers public_list avec message
        self.client.login(username="student_user", password="password123")
        response = self.client.get(reverse("sponsors:dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("sponsors:public_list"))

        # Utilisateur staff -> 200 OK
        self.client.login(username="staff_user", password="password123")
        response = self.client.get(reverse("sponsors:dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Sponsorship campaign")
        self.assertContains(response, "5,000")

    def test_sponsorship_change_status_records_history(self):
        """Le changement de statut doit mettre à jour le statut et créer une entrée StatusHistory."""
        self.client.login(username="staff_user", password="password123")

        url = reverse("sponsors:sponsorship_change_status", kwargs={"pk": self.sponsorship.pk})
        response = self.client.post(url, {"status": StatusChoices.NEGOCIATION})

        self.sponsorship.refresh_from_db()
        self.assertEqual(self.sponsorship.status, StatusChoices.NEGOCIATION)

        # Vérifier qu'un historique a été enregistré
        history = StatusHistory.objects.filter(sponsorship=self.sponsorship).latest("changed_at")
        self.assertEqual(history.old_status, StatusChoices.CONFIRME)
        self.assertEqual(history.new_status, StatusChoices.NEGOCIATION)
        self.assertEqual(history.changed_by, self.staff_user)

    def test_create_sponsor_form(self):
        """Un staff peut créer un nouveau sponsor."""
        self.client.login(username="staff_user", password="password123")
        response = self.client.post(
            reverse("sponsors:sponsor_create"),
            {
                "name": "Beta Technologies",
                "sector": "finance",
                "website": "https://beta.tn",
                "description": "Fintech innovante.",
                "contact_name": "Ines",
                "contact_email": "ines@beta.tn",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Sponsor.objects.filter(name="Beta Technologies").exists())

    def test_non_staff_post_change_status_is_rejected(self):
        """Un non-staff qui fait un POST sur le changement de statut est refusé (redirection 302, pas de modif)."""
        self.client.login(username="student_user", password="password123")
        url = reverse("sponsors:sponsorship_change_status", kwargs={"pk": self.sponsorship.pk})
        response = self.client.post(url, {"status": StatusChoices.REFUSE})

        # The role permission mixin redirects denied users to the public page.
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("sponsors:public_list"))

        # Vérifier que le statut n'a PAS changé en base de données
        self.sponsorship.refresh_from_db()
        self.assertEqual(self.sponsorship.status, StatusChoices.CONFIRME)

    def test_public_page_does_not_expose_unconfirmed_sponsors_nor_confidential_fields(self):
        """La page publique n'affiche ni un sponsor non confirmé, ni amount, ni contact_email."""
        # Création d'un sponsor prospect (non confirmé)
        unconfirmed_sponsor = Sponsor.objects.create(
            name="Secret Stealth Startup",
            website="https://stealth.tn",
            sector="tech",
            description="Projet en cours de négociation.",
            contact_name="Directeur Secret",
            contact_email="private_deal@stealth.tn",
        )
        Sponsorship.objects.create(
            sponsor=unconfirmed_sponsor,
            hackathon=self.hackathon,
            tier=TierChoices.GOLD,
            amount=Decimal("7500.000"),
            status=StatusChoices.NEGOCIATION,
            notes="Négociation en cours 7500 DT",
        )

        response = self.client.get(reverse("sponsors:public_list"))
        self.assertEqual(response.status_code, 200)

        # 1. Le sponsor non confirmé ne doit pas apparaître sur la vitrine publique
        self.assertNotContains(response, "Secret Stealth Startup")
        self.assertNotContains(response, "private_deal@stealth.tn")
        self.assertNotContains(response, "7500")

        # 2. Même pour le sponsor confirmé, amount et contact_email ne doivent JAMAIS apparaître
        self.assertContains(response, "Alpha Corp")  # Confirmé -> OK affiché
        self.assertNotContains(response, "karim@alpha.example.com")  # Email contact -> confidentiel
        self.assertNotContains(response, "5000")  # Montant financier -> confidentiel

    def test_no_history_created_if_status_identical(self):
        """Aucun historique n'est créé si le nouveau statut est identique au statut actuel."""
        initial_history_count = StatusHistory.objects.filter(sponsorship=self.sponsorship).count()

        # Appel avec le même statut actuel ('confirmé')
        result = self.sponsorship.change_status(StatusChoices.CONFIRME, self.staff_user)

        self.assertIsNone(result)
        new_history_count = StatusHistory.objects.filter(sponsorship=self.sponsorship).count()
        self.assertEqual(new_history_count, initial_history_count)

    def test_home_page_sponsors_section_dynamic_and_fallback(self):
        """Vérifie l'affichage dynamique des sponsors confirmés sur home.html et le fallback."""
        # 1. Avec sponsor confirmé : 'Alpha Corp' s'affiche dans la section
        response = self.client.get(reverse("core:home"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="sponsors"')
        self.assertContains(response, "Alpha Corp")

        # 2. Sans aucun sponsor confirmé : le fallback 'Bientôt : nos partenaires' apparaît
        self.sponsorship.status = StatusChoices.NEGOCIATION
        self.sponsorship.save(update_fields=["status"])

        response = self.client.get(reverse("core:home"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Coming soon: our partners")
        self.assertNotContains(response, "Alpha Corp")

    def test_sponsorship_form_post_update_with_status_does_not_change_status_and_no_history(self):
        """POST de modification avec un autre status => statut inchangé et aucune ligne d'historique."""
        self.client.login(username="staff_user", password="password123")
        initial_history_count = StatusHistory.objects.filter(sponsorship=self.sponsorship).count()

        url = reverse("sponsors:sponsorship_update", kwargs={"pk": self.sponsorship.pk})
        response = self.client.post(
            url,
            {
                "sponsor": self.sponsor.pk,
                "hackathon": self.hackathon.pk,
                "tier": TierChoices.PLATINUM,
                "amount": "6000.000",
                "contribution_type": self.sponsorship.contribution_type,
                "notes": "Tentative de changement détourné",
                "status": StatusChoices.REFUSE,  # Doit être ignoré car exclu de SponsorshipForm
            },
        )
        self.assertEqual(response.status_code, 302)

        self.sponsorship.refresh_from_db()
        # Le montant est bien mis à jour
        self.assertEqual(self.sponsorship.amount, Decimal("6000.000"))
        # Le statut est resté CONFIRME
        self.assertEqual(self.sponsorship.status, StatusChoices.CONFIRME)
        # Aucune entrée d'historique n'a été créée
        new_history_count = StatusHistory.objects.filter(sponsorship=self.sponsorship).count()
        self.assertEqual(new_history_count, initial_history_count)

    def test_sponsor_list_non_staff_rejected(self):
        """Un non-staff essayant d'accéder à la liste des sponsors est refusé."""
        self.client.login(username="student_user", password="password123")
        response = self.client.get(reverse("sponsors:sponsor_list"))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("sponsors:public_list"))

    def test_sponsor_list_filter_by_status_and_search_by_name(self):
        """Filtre par statut et recherche textuelle par nom dans la liste staff."""
        self.client.login(username="staff_user", password="password123")

        # Créer un deuxième sponsor
        other_sponsor = Sponsor.objects.create(name="Zeta Logistics", sector="industrie")
        Sponsorship.objects.create(
            sponsor=other_sponsor,
            hackathon=self.hackathon,
            tier=TierChoices.SILVER,
            amount=Decimal("1500.000"),
            status=StatusChoices.REFUSE,
        )

        # 1. Recherche par nom 'Alpha'
        response = self.client.get(reverse("sponsors:sponsor_list"), {"q": "Alpha"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Alpha Corp")
        self.assertNotContains(response, "Zeta Logistics")

        # 2. Filtre par statut 'refusé'
        response = self.client.get(reverse("sponsors:sponsor_list"), {"status": StatusChoices.REFUSE})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Zeta Logistics")
        self.assertNotContains(response, "Alpha Corp")

    def test_sponsor_list_pagination_and_htmx_partial(self):
        """Vérifie la pagination à 10 par page et le renvoi strict du partial sous requête HTMX."""
        self.client.login(username="staff_user", password="password123")

        # Créer 11 partenariats supplémentaires (total 12)
        for i in range(11):
            sp = Sponsor.objects.create(name=f"Company {i:02d}")
            Sponsorship.objects.create(
                sponsor=sp,
                hackathon=self.hackathon,
                tier=TierChoices.BRONZE,
                amount=Decimal("1000.000"),
                status=StatusChoices.PROSPECT,
            )

        # Requête normale : page 1 contient 10 éléments et le squelette HTML complet
        response = self.client.get(reverse("sponsors:sponsor_list"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["sponsorships"]), 10)
        self.assertTrue(response.context["is_paginated"])
        self.assertContains(response, "<!DOCTYPE html>")

        # Requête HTMX : ne doit renvoyer QUE le partial du tableau (pas de DOCTYPE)
        htmx_response = self.client.get(reverse("sponsors:sponsor_list"), HTTP_HX_REQUEST="true")
        self.assertEqual(htmx_response.status_code, 200)
        self.assertTemplateUsed(htmx_response, "sponsors/partials/sponsors_table.html")
        self.assertNotContains(htmx_response, "<!DOCTYPE html>")

    def test_get_confirmed_sponsors_deduplication(self):
        """Un sponsor confirmé sur 2 hackathons différents n'apparaît qu'une seule fois via get_confirmed_sponsors."""
        # Deuxième hackathon
        hackathon_2 = Hackathon.objects.create(
            name="ESPRIT AI Summit 2026",
            theme="Generative AI",
            start_date=timezone.now() + timezone.timedelta(days=30),
            end_date=timezone.now() + timezone.timedelta(days=32),
        )
        # Sponsoring confirmé pour le même sponsor (Alpha Corp) sur le 2e hackathon
        Sponsorship.objects.create(
            sponsor=self.sponsor,
            hackathon=hackathon_2,
            tier=TierChoices.GOLD,
            amount=Decimal("3000.000"),
            status=StatusChoices.CONFIRME,
        )

        # Deuxième sponsor distinct également confirmé
        sponsor_2 = Sponsor.objects.create(name="Beta Solutions", sector="telecom")
        Sponsorship.objects.create(
            sponsor=sponsor_2,
            hackathon=hackathon_2,
            tier=TierChoices.SILVER,
            amount=Decimal("2000.000"),
            status=StatusChoices.CONFIRME,
        )

        results = get_confirmed_sponsors(limit=6)
        # Doit contenir exactement 2 éléments (Alpha Corp et Beta Solutions), sans doublon pour Alpha Corp
        self.assertEqual(len(results), 2)
        sponsor_ids = [item.sponsor_id for item in results]
        self.assertEqual(len(sponsor_ids), len(set(sponsor_ids)))
        self.assertIn(self.sponsor.id, sponsor_ids)
        self.assertIn(sponsor_2.id, sponsor_ids)

    def test_sponsor_list_no_n_plus_one_queries(self):
        """Vérifie que la liste staff utilise select_related et évite les requêtes N+1."""
        self.client.login(username="staff_user", password="password123")

        # Créer plusieurs partenariats pour tester le volume
        for i in range(5):
            sp = Sponsor.objects.create(name=f"Enterprise {i}")
            Sponsorship.objects.create(
                sponsor=sp,
                hackathon=self.hackathon,
                tier=TierChoices.BRONZE,
                amount=Decimal("1000.000"),
                status=StatusChoices.CONFIRME,
            )

        # Le nombre de requêtes SQL doit être borné et constant (aucun problème N+1)
        # 1: Session, 2: Auth User, 3: Count pagination, 4: User Profile navbar, 5: Hackathons dropdown, 6: Sponsorships + select_related(sponsor, hackathon)
        with self.assertNumQueries(6):
            response = self.client.get(reverse("sponsors:sponsor_list"))
            self.assertEqual(response.status_code, 200)

    def test_sponsor_contact_non_staff_rejected(self):
        """Un non-staff essayant d'accéder ou d'envoyer un email à un sponsor est refusé."""
        # 1. Anonyme GET -> Redirigé vers login
        url = reverse("sponsors:sponsor_contact", kwargs={"pk": self.sponsor.pk})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith("/accounts/login/"))

        # 2. Utilisateur non-staff GET -> Redirigé vers public_list
        self.client.login(username="student_user", password="password123")
        response = self.client.get(url)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("sponsors:public_list"))

        # 3. Utilisateur non-staff POST -> Refusé, aucun email envoyé
        response = self.client.post(
            url,
            {"subject": "Hack attempt", "message": "Trying to send unauthorized email."},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(len(mail.outbox), 0)
        self.assertEqual(SponsorEmailLog.objects.count(), 0)

    def test_sponsor_contact_send_success_and_logging(self):
        """Un organisateur staff peut envoyer un email qui est tracé dans SponsorEmailLog."""
        self.client.login(username="staff_user", password="password123")
        url = reverse("sponsors:sponsor_contact", kwargs={"pk": self.sponsor.pk})

        # 1. GET affiche le formulaire avec modèle prérempli
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Contact Alpha Corp")
        self.assertContains(response, self.sponsor.contact_email)

        # 2. POST envoi d'email
        post_data = {
            "template_key": "invitation",
            "subject": "Invitation Partenariat ESPRIT Hack 2026",
            "message": "Bonjour Karim,\nNous vous invitons à devenir partenaire officiel.",
        }
        post_response = self.client.post(url, post_data)
        self.assertEqual(post_response.status_code, 302)
        self.assertRedirects(post_response, reverse("sponsors:sponsor_detail", kwargs={"pk": self.sponsor.pk}))

        # Vérifier l'envoi effectif via outbox Django
        self.assertEqual(len(mail.outbox), 1)
        sent_email = mail.outbox[0]
        self.assertEqual(sent_email.to, [self.sponsor.contact_email])
        self.assertEqual(sent_email.subject, "Invitation Partenariat ESPRIT Hack 2026")
        self.assertIn("Bonjour Karim", sent_email.body)

        # Vérifier la création de la trace SponsorEmailLog
        log = SponsorEmailLog.objects.filter(sponsor=self.sponsor).latest("sent_at")
        self.assertTrue(log.success)
        self.assertEqual(log.sent_by, self.staff_user)
        self.assertEqual(log.subject, "Invitation Partenariat ESPRIT Hack 2026")
        self.assertEqual(log.error_message, "")

    def test_sponsor_contact_recipient_cannot_be_overridden_via_post(self):
        """Le destinataire est STRICTEMENT sponsor.contact_email et ne peut pas être manipulé en POST."""
        self.client.login(username="staff_user", password="password123")
        url = reverse("sponsors:sponsor_contact", kwargs={"pk": self.sponsor.pk})

        # Tentative d'injection d'une autre adresse email dans le payload POST
        post_data = {
            "template_key": "relance",
            "subject": "Test Sujet Sécurité",
            "message": "Contenu du test sécurité.",
            "recipient": "evil_hacker@darkweb.org",
            "contact_email": "evil_hacker@darkweb.org",
            "to": "evil_hacker@darkweb.org",
        }
        response = self.client.post(url, post_data)
        self.assertEqual(response.status_code, 302)

        # L'email doit impérativement avoir été envoyé au contact légitime
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, [self.sponsor.contact_email])
        self.assertNotIn("evil_hacker@darkweb.org", mail.outbox[0].to)

    def test_sponsor_contact_rate_limit_exceeded(self):
        """La limite anti-abus bloque après 5 emails par minute pour le même utilisateur."""
        cache.clear()
        self.client.login(username="staff_user", password="password123")
        url = reverse("sponsors:sponsor_contact", kwargs={"pk": self.sponsor.pk})

        # Envoyer 5 emails autorisés
        for i in range(5):
            res = self.client.post(url, {"subject": f"Email #{i+1}", "message": f"Message #{i+1}"})
            self.assertEqual(res.status_code, 302)

        self.assertEqual(len(mail.outbox), 5)

        # Le 6ème envoi dans la même minute doit être rejeté par la limite
        res_6 = self.client.post(url, {"subject": "Email #6", "message": "Trop de messages"})
        self.assertEqual(res_6.status_code, 200)  # Re-rendu de la page avec message d'erreur
        self.assertContains(res_6, "Rate limit exceeded")

        # Pas de 6ème email envoyé
        self.assertEqual(len(mail.outbox), 5)

        # Trace enregistrée avec success=False
        last_log = SponsorEmailLog.objects.filter(sponsor=self.sponsor).latest("sent_at")
        self.assertFalse(last_log.success)
        self.assertIn("Rate limit", last_log.error_message)

    def test_sponsor_contact_handles_failure_gracefully(self):
        """En cas d'échec d'envoi réseau/SMTP, l'application ne crashe pas et enregistre l'erreur."""
        cache.clear()
        self.client.login(username="staff_user", password="password123")
        url = reverse("sponsors:sponsor_contact", kwargs={"pk": self.sponsor.pk})

        with patch("sponsors.services.send_mail", side_effect=Exception("SMTP Connection Refused: port 587")):
            response = self.client.post(
                url,
                {"subject": "Objet test panne", "message": "Message lors d'une panne SMTP."},
            )

        # La page ne doit pas crasher (erreur 500 évitée, formulaire réaffiché avec message d'erreur)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "SMTP Connection Refused")

        # Un log a bien été créé avec success=False et le message d'erreur
        failed_log = SponsorEmailLog.objects.filter(sponsor=self.sponsor).latest("sent_at")
        self.assertFalse(failed_log.success)
        self.assertIn("SMTP Connection Refused", failed_log.error_message)

    # ==========================================================================
    # TESTS ÉTAPE 8 : RECOMMANDATION VECTORIELLE (TF-IDF + COSINUS)
    # ==========================================================================

    def test_ai_recommendation_closest_sector_and_theme_ranks_first(self):
        """Le sponsor ayant la plus forte proximité sémantique avec le thème du hackathon arrive en tête."""
        cache.clear()

        # Hackathon centré sur la fintech et la blockchain
        fintech_hackathon = Hackathon.objects.create(
            name="FinTech Innovate ESPRIT",
            theme="Finance digitale, paiements mobiles et blockchain bancaire",
            start_date=timezone.now(),
            end_date=timezone.now() + timezone.timedelta(days=2),
        )

        # Sponsor 1 : Fintech & Banque (forte affinité)
        fintech_sponsor = Sponsor.objects.create(
            name="PayFlow Tech",
            sector="finance",
            description="Solutions de paiement mobile et d'infrastructure pour la finance digitale.",
        )

        # Sponsor 2 : Agro-alimentaire (faible affinité)
        agro_sponsor = Sponsor.objects.create(
            name="AgriBio Food",
            sector="sante",
            description="Production agricole biologique et élevage durable.",
        )

        # Sponsor 3 : Énergie (affinité neutre)
        energy_sponsor = Sponsor.objects.create(
            name="SolarGrid Energy",
            sector="energie",
            description="Panneaux solaires et réseaux électriques intelligents.",
        )

        data = recommend_sponsors_for_hackathon(fintech_hackathon, use_cache=False)
        recs = data["recommendations"]

        self.assertGreaterEqual(len(recs), 2)
        # PayFlow Tech doit être premier du classement
        self.assertEqual(recs[0]["sponsor"], fintech_sponsor)
        self.assertGreater(recs[0]["score"], recs[-1]["score"])

        # Vérifier que des mots-clés explicatifs pertinents sont extraits
        keywords = recs[0]["keywords"]
        self.assertGreater(len(keywords), 0)
        self.assertTrue(any(kw in ["finance", "paiement", "digitale", "mobile", "tech", "payflow"] for kw in keywords))

    def test_ai_recommendation_excludes_existing_partnerships_for_same_hackathon(self):
        """Les sponsors ayant déjà un partenariat sur ce hackathon sont strictement exclus des suggestions."""
        cache.clear()
        # self.sponsor a déjà un partenariat sur self.hackathon dans setUp()
        # Création de deux sponsors disponibles
        sp1 = Sponsor.objects.create(name="Available One", sector="tech")
        sp2 = Sponsor.objects.create(name="Available Two", sector="telecom")

        data = recommend_sponsors_for_hackathon(self.hackathon, use_cache=False)
        rec_sponsor_ids = [r["sponsor"].id for r in data["recommendations"]]

        # self.sponsor (Alpha Corp) NE doit PAS figurer dans les recommandations
        self.assertNotIn(self.sponsor.id, rec_sponsor_ids)
        self.assertIn(sp1.id, rec_sponsor_ids)
        self.assertIn(sp2.id, rec_sponsor_ids)

    def test_ai_recommendation_insufficient_data_handled_gracefully(self):
        """Le cas 'pas assez de données' (moins de 2 sponsors ou texte vide) est géré proprement sans erreur."""
        cache.clear()
        empty_hackathon = Hackathon.objects.create(name="", theme="")

        data = recommend_sponsors_for_hackathon(empty_hackathon, use_cache=False)
        self.assertEqual(data["recommendations"], [])
        self.assertIn("insufficient", data["message"].lower())

    def test_ai_recommendations_view_non_staff_rejected(self):
        """L'accès à l'endpoint de recommandations requiert les privilèges Staff."""
        url = reverse("sponsors:sponsor_recommendations", kwargs={"hackathon_id": self.hackathon.pk})

        # 1. Anonyme -> 302
        response = self.client.get(url)
        self.assertEqual(response.status_code, 302)

        # 2. Non-staff -> 302 vers public_list
        self.client.login(username="student_user", password="password123")
        response = self.client.get(url)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("sponsors:public_list"))

        # 3. Staff -> 200 OK
        self.client.login(username="staff_user", password="password123")
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)

    # ==========================================================================
    # TESTS ÉTAPE 8 : GÉNÉRATION D'EMAIL PAR IA (LLM & FALLBACK)
    # ==========================================================================

    def test_ai_email_generation_without_api_key_falls_back_to_template(self):
        """Sans clé API configurée, le système bascule automatiquement sur les modèles standard."""
        self.client.login(username="staff_user", password="password123")
        url = reverse("sponsors:sponsor_ai_generate_email", kwargs={"pk": self.sponsor.pk})

        with patch("os.getenv", return_value=""):
            response = self.client.post(
                url,
                {"template_key": "invitation", "tone": "professionnel"},
            )

        self.assertEqual(response.status_code, 200)
        # Vérifier que le corps correspond au modèle standard d'invitation
        self.assertContains(response, "Invitation to support")
        # Notice explicative
        self.assertContains(response, "API key is not configured")

    def test_ai_email_generation_with_mocked_api_success(self):
        """Avec une clé API et une réponse API valide mockée, le texte IA est injecté."""
        self.client.login(username="staff_user", password="password123")
        url = reverse("sponsors:sponsor_ai_generate_email", kwargs={"pk": self.sponsor.pk})

        # Groq uses the OpenAI-compatible response shape.
        fake_api_response = json.dumps({
            "choices": [
                {
                    "message": {
                        "content": (
                            "SUBJECT: Excellence in partnership with Alpha Corp — ESPRIT Hack 2026\n\n"
                            "Dear Karim,\nJoin Hackathon 2026."
                        )
                    }
                }
            ]
        }).encode("utf-8")

        mock_resp = MagicMock()
        mock_resp.read.return_value = fake_api_response
        mock_resp.__enter__.return_value = mock_resp

        with patch.dict(os.environ, {"GROQ_API_KEY": "gsk_test_key_mock"}):
            with patch("urllib.request.urlopen", return_value=mock_resp):
                response = self.client.post(
                    url,
                    {"template_key": "invitation", "tone": "professionnel"},
                )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Excellence in partnership with Alpha Corp")
        self.assertContains(response, "Join Hackathon 2026.")
        self.assertContains(response, "AI enabled")

    def test_ai_email_generation_api_error_falls_back_gracefully(self):
        """En cas d'erreur réseau ou timeout avec le fournisseur IA, repli immédiat sans crash."""
        self.client.login(username="staff_user", password="password123")
        url = reverse("sponsors:sponsor_ai_generate_email", kwargs={"pk": self.sponsor.pk})

        with patch.dict(os.environ, {"GROQ_API_KEY": "gsk_test_key_mock"}):
            with patch(
                "urllib.request.urlopen",
                side_effect=Exception("HTTP 504 Gateway Timeout"),
            ):
                response = self.client.post(
                    url,
                    {"template_key": "relance", "tone": "amical"},
                )

        self.assertEqual(response.status_code, 200)
        # Repli sur le modèle relance
        self.assertContains(response, "Following up on our sponsorship proposal")
        self.assertContains(response, "AI is temporarily unavailable")

    def test_ai_email_generation_payload_never_leaks_confidential_fields(self):
        """Le payload envoyé à l'IA ne contient JAMAIS amount, notes, contact_name ni contact_email."""
        confidential_sponsor = Sponsor.objects.create(
            name="Confidential Bank",
            sector="finance",
            description="Banque privée d'investissement.",
            contact_name="Directeur Secret",
            contact_email="private_ceo@bank.tn",
        )
        Sponsorship.objects.create(
            sponsor=confidential_sponsor,
            hackathon=self.hackathon,
            tier=TierChoices.PLATINUM,
            amount=Decimal("45000.000"),
            notes="Dossier strictement confidentiel - Négociation exclusive 45000 DT.",
        )

        payload = build_sanitized_ai_payload(confidential_sponsor, self.hackathon, tone="professionnel")
        serialized_payload = str(payload)

        # 1. Vérifier la présence des informations publiques autorisées
        self.assertIn("Confidential Bank", serialized_payload)
        self.assertIn("Banque privée d'investissement", serialized_payload)

        # 2. Vérifier l'absence absolue des données confidentielles
        self.assertNotIn("45000", serialized_payload)
        self.assertNotIn("Dossier strictement confidentiel", serialized_payload)
        self.assertNotIn("Directeur Secret", serialized_payload)
        self.assertNotIn("private_ceo@bank.tn", serialized_payload)

    # ==========================================================================
    # TESTS ÉTAPE 6 : EXPORT EXCEL DES PARTENARIATS (openpyxl)
    # ==========================================================================

    def test_excel_export_non_staff_rejected(self):
        """L'export Excel est strictement réservé aux utilisateurs Staff."""
        url = reverse("sponsors:sponsor_export_excel")

        # 1. Anonyme -> 302 vers login
        response = self.client.get(url)
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith("/accounts/login/"))

        # 2. Non-staff -> 302 vers public_list
        self.client.login(username="student_user", password="password123")
        response = self.client.get(url)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("sponsors:public_list"))

    def test_excel_export_success_and_valid_openpyxl_structure(self):
        """Un staff peut télécharger un classeur Excel valide avec deux feuilles (Partenariats + Résumé)."""
        self.client.login(username="staff_user", password="password123")
        url = reverse("sponsors:sponsor_export_excel")

        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("hackflow_sponsors_", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])

        # Ouverture et vérification du contenu avec openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["Sponsorships", "Summary"])

        # Vérifier la feuille 1
        ws_data = wb["Sponsorships"]
        headers = [cell.value for cell in ws_data[1]]
        self.assertIn("Sponsor", headers)
        self.assertIn("Amount (DT)", headers)
        self.assertIn("Hackathon", headers)

        # Vérifier la ligne de données
        row_2_values = [cell.value for cell in ws_data[2]]
        self.assertIn("Alpha Corp", row_2_values)
        self.assertIn("ESPRIT Hack 2026", row_2_values)
        self.assertIn(5000.0, row_2_values)

        # Vérifier la feuille 2
        ws_summary = wb["Summary"]
        self.assertIn("Financial and operational sponsorship summary", ws_summary["A1"].value)

    def test_excel_export_respects_active_filters(self):
        """L'export Excel applique fidèlement les filtres passés en paramètres GET."""
        self.client.login(username="staff_user", password="password123")

        # Sponsor refusé
        other_sponsor = Sponsor.objects.create(name="Excluded Company", sector="finance")
        Sponsorship.objects.create(
            sponsor=other_sponsor,
            hackathon=self.hackathon,
            tier=TierChoices.SILVER,
            amount=Decimal("1500.000"),
            status=StatusChoices.REFUSE,
        )

        # Export filtré uniquement sur status=confirmé
        url = reverse("sponsors:sponsor_export_excel") + "?status=confirmé"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws_data = wb["Sponsorships"]

        sponsor_names = [ws_data.cell(row=r, column=1).value for r in range(2, ws_data.max_row + 1)]
        self.assertIn("Alpha Corp", sponsor_names)
        self.assertNotIn("Excluded Company", sponsor_names)

    def test_excel_export_formula_injection_neutralized(self):
        """Toute valeur débutant par =, +, -, @ est neutralisée pour bloquer l'injection de formules."""
        self.client.login(username="staff_user", password="password123")

        # Sponsor avec nom malveillant
        malicious_sponsor = Sponsor.objects.create(
            name="=1+1",
            sector="tech",
            description="+SUM(A1:A10)",
        )
        Sponsorship.objects.create(
            sponsor=malicious_sponsor,
            hackathon=self.hackathon,
            tier=TierChoices.BRONZE,
            amount=Decimal("1000.000"),
            status=StatusChoices.PROSPECT,
        )

        response = self.client.get(reverse("sponsors:sponsor_export_excel"))
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws_data = wb["Sponsorships"]

        # Recherche de la ligne du sponsor malveillant et vérification de la neutralisation
        found_cell = False
        for r in range(2, ws_data.max_row + 1):
            cell = ws_data.cell(row=r, column=1)
            val = str(cell.value or "")
            if val == "=1+1":
                # La cellule ne doit pas avoir d'apostrophe visible, mais son type doit être strictement texte ("s")
                self.assertEqual(cell.data_type, "s")
                self.assertFalse(val.startswith("'"))
                found_cell = True
        self.assertTrue(found_cell, "La cellule avec formule aurait dû être trouvée et neutralisée avec le type texte ('s').")

    def test_excel_export_row_limit_handled_gracefully(self):
        """Au-delà de 5000 lignes, l'export affiche un message clair sans planter."""
        self.client.login(username="staff_user", password="password123")
        url = reverse("sponsors:sponsor_export_excel")

        # Mock de count() retournant 5001
        with patch.object(Sponsorship.objects, "count", return_value=5001):
            with patch("sponsors.views.filter_sponsorships_queryset") as mock_filter:
                mock_qs = MagicMock()
                mock_qs.count.return_value = 5001
                mock_filter.return_value = mock_qs

                response = self.client.get(url)

        # Redirigé vers la liste avec message d'erreur
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("sponsors:sponsor_list"))


class SeedSponsorsCommandTest(TestCase):
    def setUp(self):
        self.staff_user = User.objects.create_user(
            username="seed_test_staff",
            password="test-password",
            is_staff=True,
        )
        self.legacy_sponsor = Sponsor.objects.create(
            name="InstaDeep Labs",
            website="https://instadeep.com",
            sector="tech",
            description="Existing record that must remain untouched.",
        )
        self.legacy_hackathon = Hackathon.objects.create(
            name="Existing Hackathon",
            theme="Existing event",
            start_date=timezone.now().date(),
            end_date=(timezone.now() + timezone.timedelta(days=2)).date(),
        )
        self.legacy_sponsorship = Sponsorship.objects.create(
            sponsor=self.legacy_sponsor,
            hackathon=self.legacy_hackathon,
            tier=TierChoices.GOLD,
            amount=Decimal("1000.000"),
            status=StatusChoices.PROSPECT,
        )

    def test_seed_does_not_update_sponsors_without_demo_suffix(self):
        call_command("seed_sponsors", force=True, stdout=io.StringIO())

        self.legacy_sponsor.refresh_from_db()
        self.legacy_sponsorship.refresh_from_db()
        self.assertEqual(self.legacy_sponsor.website, "https://instadeep.com")
        self.assertEqual(
            self.legacy_sponsor.description,
            "Existing record that must remain untouched.",
        )
        self.assertEqual(self.legacy_sponsorship.status, StatusChoices.PROSPECT)
        self.assertTrue(Sponsor.objects.filter(name="Northstar AI (demo)").exists())

    def test_reset_preserves_non_demo_records_and_recreates_demo_data(self):
        call_command("seed_sponsors", force=True, stdout=io.StringIO())
        demo_hackathon = Hackathon.objects.get(name="HackFlow Challenge (demo)")
        legacy_demo_event_partnership = Sponsorship.objects.create(
            sponsor=self.legacy_sponsor,
            hackathon=demo_hackathon,
            tier=TierChoices.SILVER,
            amount=Decimal("2000.000"),
            status=StatusChoices.CONTACTE,
        )

        call_command(
            "seed_sponsors",
            force=True,
            reset=True,
            stdout=io.StringIO(),
        )

        self.legacy_sponsor.refresh_from_db()
        self.legacy_sponsorship.refresh_from_db()
        legacy_demo_event_partnership.refresh_from_db()
        self.assertEqual(self.legacy_sponsor.website, "https://instadeep.com")
        self.assertEqual(self.legacy_sponsorship.status, StatusChoices.PROSPECT)
        self.assertEqual(
            legacy_demo_event_partnership.status,
            StatusChoices.CONTACTE,
        )
        self.assertEqual(
            Sponsor.objects.filter(name__endswith="(demo)").count(),
            5,
        )
        self.assertEqual(
            Sponsorship.objects.filter(
                hackathon=demo_hackathon,
                sponsor__name__endswith="(demo)",
            ).count(),
            5,
        )


class SponsorRolePermissionTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="role_matrix_user",
            password="test-password",
        )
        self.staff_user = User.objects.create_user(
            username="role_matrix_staff",
            password="test-password",
            is_staff=True,
        )
        self.sponsor = Sponsor.objects.create(
            name="Role Matrix Sponsor",
            sector="tech",
            contact_email="contact@example.com",
        )
        self.other_sponsor = Sponsor.objects.create(
            name="Role Matrix Candidate",
            sector="tech",
            description="AI and cloud",
        )
        self.hackathon = Hackathon.objects.create(
            name="Role Matrix Hackathon",
            theme="AI and cloud",
            start_date=timezone.now().date(),
            end_date=(timezone.now() + timezone.timedelta(days=2)).date(),
        )
        self.sponsorship = Sponsorship.objects.create(
            sponsor=self.sponsor,
            hackathon=self.hackathon,
            tier=TierChoices.GOLD,
            amount=Decimal("1200.000"),
            status=StatusChoices.PROSPECT,
        )

    def test_default_role_permissions(self):
        expected = {
            "admin": SPONSOR_PERMISSIONS,
            "organisateur": {"manage", "view_confidential", "contact", "export"},
            "etudiant": set(),
            "jury": set(),
            "visiteur": set(),
        }
        self.assertEqual(ROLE_PERMISSIONS, expected)

        for role, permissions in expected.items():
            for permission in SPONSOR_PERMISSIONS:
                with self.subTest(role=role, permission=permission):
                    with patch(
                        "sponsors.permissions.get_user_roles",
                        return_value={role},
                    ):
                        self.assertEqual(
                            user_has_sponsor_permission(self.user, permission),
                            permission in permissions,
                        )

    def test_each_role_permission_controls_its_views(self):
        routes_by_permission = {
            "manage": [
                ("get", "sponsors:sponsor_create", {}),
                (
                    "get",
                    "sponsors:sponsor_update",
                    {"kwargs": {"pk": self.sponsor.pk}},
                ),
                ("get", "sponsors:sponsorship_create", {}),
                (
                    "get",
                    "sponsors:sponsorship_update",
                    {"kwargs": {"pk": self.sponsorship.pk}},
                ),
                (
                    "post",
                    "sponsors:sponsorship_change_status",
                    {
                        "kwargs": {"pk": self.sponsorship.pk},
                        "data": {"status": StatusChoices.NEGOCIATION},
                    },
                ),
            ],
            "view_confidential": [
                ("get", "sponsors:dashboard", {}),
                ("get", "sponsors:sponsor_list", {}),
                (
                    "get",
                    "sponsors:sponsor_detail",
                    {"kwargs": {"pk": self.sponsor.pk}},
                ),
                (
                    "get",
                    "sponsors:sponsor_recommendations",
                    {"kwargs": {"hackathon_id": self.hackathon.pk}},
                ),
            ],
            "contact": [
                (
                    "get",
                    "sponsors:sponsor_contact",
                    {"kwargs": {"pk": self.sponsor.pk}},
                ),
                (
                    "post",
                    "sponsors:sponsor_ai_generate_email",
                    {
                        "kwargs": {"pk": self.sponsor.pk},
                        "data": {"template_key": "invitation"},
                    },
                ),
            ],
            "export": [("get", "sponsors:sponsor_export_excel", {})],
            "delete": [
                (
                    "get",
                    "sponsors:sponsor_delete",
                    {"kwargs": {"pk": self.sponsor.pk}},
                ),
                (
                    "get",
                    "sponsors:sponsorship_delete",
                    {"kwargs": {"pk": self.sponsorship.pk}},
                ),
            ],
        }

        for role, permissions in ROLE_PERMISSIONS.items():
            self.client.force_login(self.user)
            with patch(
                "sponsors.permissions.get_user_roles",
                return_value={role},
            ):
                for permission, routes in routes_by_permission.items():
                    expected_access = permission in permissions
                    for method, route_name, options in routes:
                        with self.subTest(
                            role=role,
                            permission=permission,
                            route=route_name,
                        ):
                            url = reverse(
                                route_name,
                                kwargs=options.get("kwargs"),
                            )
                            request_method = getattr(self.client, method)
                            if route_name == "sponsors:sponsor_ai_generate_email":
                                with patch(
                                    "sponsors.views.generate_prospecting_email",
                                    return_value={
                                        "subject": "Demo",
                                        "body": "Demo message",
                                        "ai_generated": False,
                                    },
                                ):
                                    response = request_method(
                                        url,
                                        options.get("data", {}),
                                    )
                            elif method == "post":
                                response = request_method(
                                    url,
                                    options.get("data", {}),
                                )
                            else:
                                response = request_method(url)
                            self.assertEqual(
                                response.status_code == 200
                                or response.status_code == 302
                                and response.url != reverse("sponsors:public_list"),
                                expected_access,
                            )
                            if not expected_access:
                                self.assertEqual(response.status_code, 302)
                                self.assertEqual(
                                    response.url,
                                    reverse("sponsors:public_list"),
                                )

    @override_settings(SPONSORS_USE_PROFILE_ROLE=True)
    def test_profile_role_source_works_without_staff_flag(self):
        # Patch the saved instance, not the model class: `role` is a field
        # descriptor, so patching the class never changes user.profile.role.
        self.user.profile.role = "organizer"  # accounts.Profile.Role.ORGANIZER
        self.user.profile.save(update_fields=["role"])

        self.assertFalse(self.user.is_staff)
        # accounts roles are mapped onto the French label used here.
        self.assertEqual(get_user_roles(self.user), {"organisateur"})
        self.assertTrue(can_view_confidential(self.user))
        self.assertTrue(can_manage_sponsorship(self.user, self.sponsorship))
        self.assertEqual(
            sponsorships_visible_to(self.user).count(),
            Sponsorship.objects.count(),
        )

        self.client.force_login(self.user)
        response = self.client.get(reverse("sponsors:dashboard"))
        self.assertEqual(response.status_code, 200)

    @override_settings(SPONSORS_ROLE_PERMISSIONS={"jury": {"manage"}})
    def test_role_permission_map_can_be_overridden_in_settings(self):
        with patch(
            "sponsors.permissions.get_user_roles",
            return_value={"jury"},
        ):
            self.assertTrue(user_has_sponsor_permission(self.user, "manage"))
            self.assertFalse(user_has_sponsor_permission(self.user, "delete"))
